from __future__ import annotations

import textwrap
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "documentacion"
IMG_DIR = OUT_DIR / "img"
DOCX_PATH = OUT_DIR / "Flujo_Cuadre_Automatico_Original_vs_Nuevo.docx"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
GREEN = "E8F5E9"
GOLD = "FFF8E1"
RED = "FDECEC"
BORDER = "B8C7D9"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def draw_wrapped_text(draw: ImageDraw.ImageDraw, box, text, fnt, fill=(15, 35, 69), align="center"):
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 24
    words = text.split()
    lines = []
    current = ""
    for word in words:
        probe = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), probe, font=fnt)[2] <= max_width:
            current = probe
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    line_h = draw.textbbox((0, 0), "Ag", font=fnt)[3] + 8
    total_h = line_h * len(lines)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        if align == "center":
            x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2
        else:
            x = x1 + 14
        draw.text((x, y), line, font=fnt, fill=fill)
        y += line_h


def arrow(draw: ImageDraw.ImageDraw, start, end, fill=(31, 77, 120), width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) > abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        pts = [(x2, y2), (x2 - 16 * direction, y2 - 9), (x2 - 16 * direction, y2 + 9)]
    else:
        direction = 1 if y2 > y1 else -1
        pts = [(x2, y2), (x2 - 9, y2 - 16 * direction), (x2 + 9, y2 - 16 * direction)]
    draw.polygon(pts, fill=fill)


def rounded_box(draw, xy, fill, outline=BORDER, radius=18, width=2):
    if isinstance(fill, str) and not fill.startswith("#"):
        fill = "#" + fill
    if isinstance(outline, str) and not outline.startswith("#"):
        outline = "#" + outline
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def pil_color(value):
    if isinstance(value, str) and not value.startswith("#"):
        return "#" + value
    return value


def make_legacy_diagram(path: Path):
    img = Image.new("RGB", (1500, 880), "white")
    d = ImageDraw.Draw(img)
    title_f = font(36, True)
    box_f = font(24, True)
    small_f = font(20)
    d.text((50, 35), "Flujo original legacy - Cuadre Automatico WinForms", font=title_f, fill=(11, 37, 69))

    boxes = [
        ((70, 120, 360, 230), "Usuario ejecuta formulario legacy"),
        ((470, 120, 790, 230), "Consulta rutas no cuadradas\nsp_TraerRutasNoCuadradas"),
        ((900, 120, 1250, 230), "Itera ruta por ruta"),
        ((150, 325, 500, 465), "Prepara datos:\nsaldo + productos + retiros + ventas"),
        ((610, 325, 960, 465), "Valida:\ncuadre existente, movimientos,\nsaldos y ventas"),
        ((1070, 325, 1420, 465), "Registra cuadre:\ntbl_Cuadre + tbl_CodigoCuadre"),
        ((380, 600, 720, 730), "Registra retiros\ntbl_SaldoRetiro"),
        ((820, 600, 1160, 730), "Actualiza saldo ruta\ny equipo tecnico"),
    ]
    for i, (xy, text) in enumerate(boxes):
        fill = LIGHT_BLUE if i < 3 else (GREEN if i in (5, 6, 7) else LIGHT_GRAY)
        rounded_box(d, xy, fill)
        draw_wrapped_text(d, xy, text, box_f if i < 3 else small_f)

    arrow(d, (360, 175), (470, 175))
    arrow(d, (790, 175), (900, 175))
    arrow(d, (1075, 230), (325, 325))
    arrow(d, (500, 395), (610, 395))
    arrow(d, (960, 395), (1070, 395))
    arrow(d, (1245, 465), (720, 600))
    arrow(d, (720, 665), (820, 665))

    note = (
        "Formula principal: ItemsSobrantes = saldo inicial - ventas del dia. "
        "Los retiros se registran como detalle separado."
    )
    d.rounded_rectangle((70, 785, 1430, 840), radius=15, fill=pil_color(GOLD), outline=(220, 170, 60), width=2)
    draw_wrapped_text(d, (70, 785, 1430, 840), note, small_f, fill=(90, 68, 0))
    img.save(path)


def make_new_diagram(path: Path):
    img = Image.new("RGB", (1500, 930), "white")
    d = ImageDraw.Draw(img)
    title_f = font(36, True)
    box_f = font(22, True)
    small_f = font(19)
    d.text((50, 35), "Flujo nuevo - Usuario SISTEMAS en front web + backend Spring", font=title_f, fill=(11, 37, 69))

    boxes = [
        ((70, 120, 390, 235), "Login usuario\nsistemas"),
        ((500, 120, 820, 235), "Pantalla\nCuadre automatico"),
        ((930, 120, 1280, 235), "Preview:\nlista rutas pendientes"),
        ((70, 350, 390, 490), "Ejecutar:\nPOST /cuadre/sistemas/automatico/ejecutar"),
        ((500, 350, 820, 490), "Por cada ruta:\nvalidar cierre, movimientos y cuadre previo"),
        ((930, 350, 1280, 490), "Construir detalle:\nsaldo, usado hoy, retiros, sobrante"),
        ((285, 625, 635, 765), "Validar detalle:\nsobrante >= 0\nsaldo = sobrante + vendido"),
        ((785, 625, 1135, 765), "Transaccion SQL:\nregistrar cuadre, detalle,\nretiros y actualizar saldo"),
    ]
    for i, (xy, text) in enumerate(boxes):
        fill = LIGHT_BLUE if i < 3 else (GREEN if i == 7 else LIGHT_GRAY)
        rounded_box(d, xy, fill)
        draw_wrapped_text(d, xy, text, box_f if i < 3 else small_f)

    arrow(d, (390, 178), (500, 178))
    arrow(d, (820, 178), (930, 178))
    arrow(d, (1105, 235), (230, 350))
    arrow(d, (390, 420), (500, 420))
    arrow(d, (820, 420), (930, 420))
    arrow(d, (1105, 490), (460, 625))
    arrow(d, (635, 695), (785, 695))

    d.rounded_rectangle((70, 825, 1430, 885), radius=15, fill=pil_color(GOLD), outline=(220, 170, 60), width=2)
    draw_wrapped_text(
        d,
        (70, 825, 1430, 885),
        "Resultado por ruta: Registrado, Omitido o Error. El resumen alimenta contadores de la pantalla.",
        small_f,
        fill=(90, 68, 0),
    )
    img.save(path)


def make_data_diagram(path: Path):
    img = Image.new("RGB", (1500, 700), "white")
    d = ImageDraw.Draw(img)
    title_f = font(36, True)
    box_f = font(22, True)
    small_f = font(19)
    d.text((50, 35), "Modelo de calculo de saldo y registro", font=title_f, fill=(11, 37, 69))

    boxes = [
        ((80, 140, 380, 260), "tbl_saldotarjetas\nsaldo inicial por ruta/producto"),
        ((470, 140, 820, 260), "tbl_venta + tbl_codigoventa\nusado hoy\nv.e_eliminado=0\ncv.e_eliminado=0"),
        ((910, 140, 1220, 260), "Sobrante calculado\nsaldo - usado hoy"),
        ((260, 420, 560, 540), "tbl_Cuadre\ncabecera"),
        ((650, 420, 950, 540), "tbl_CodigoCuadre\ndetalle por producto"),
        ((1040, 420, 1340, 540), "tbl_SaldoRetiro\nproductos retirados"),
    ]
    fills = [LIGHT_BLUE, LIGHT_BLUE, GREEN, LIGHT_GRAY, LIGHT_GRAY, LIGHT_GRAY]
    for (xy, text), fill in zip(boxes, fills):
        rounded_box(d, xy, fill)
        draw_wrapped_text(d, xy, text, box_f if "Sobrante" in text else small_f)

    arrow(d, (380, 200), (470, 200))
    arrow(d, (820, 200), (910, 200))
    arrow(d, (1065, 260), (800, 420))
    arrow(d, (1065, 260), (410, 420))
    arrow(d, (1065, 260), (1190, 420))
    d.rounded_rectangle((80, 600, 1420, 650), radius=15, fill=pil_color(RED), outline=(180, 90, 90), width=2)
    draw_wrapped_text(
        d,
        (80, 600, 1420, 650),
        "Regla critica: si la venta o el codigo venta esta eliminado, no participa en el usado del dia.",
        small_f,
        fill=(100, 20, 20),
    )
    img.save(path)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9E2EC"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table, header_fill=LIGHT_GRAY):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        if row_idx == 0:
            tr_pr = row._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9.5)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(11, 37, 69)


def add_table(doc: Document, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].width = Inches(widths[i])
    style_table(table)
    return table


def add_bullets(doc: Document, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items):
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Paso"
    table.rows[0].cells[1].text = "Descripcion"
    table.rows[0].cells[0].width = Inches(0.7)
    table.rows[0].cells[1].width = Inches(5.7)
    for idx, item in enumerate(items, 1):
        cells = table.add_row().cells
        cells[0].text = str(idx)
        cells[1].text = item
        cells[0].width = Inches(0.7)
        cells[1].width = Inches(5.7)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_table(table)
    doc.add_paragraph()


def add_note(doc: Document, title: str, text: str, fill=GOLD):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, "E4C46A")
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor(90, 68, 0)
    p.add_run(text)
    p.paragraph_format.space_after = Pt(0)


def setup_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.167


def add_footer(doc: Document):
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Cuadre automatico - Documento tecnico")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 100, 115)


def build_doc():
    OUT_DIR.mkdir(exist_ok=True)
    IMG_DIR.mkdir(exist_ok=True)
    legacy_img = IMG_DIR / "flujo_legacy.png"
    new_img = IMG_DIR / "flujo_nuevo.png"
    data_img = IMG_DIR / "modelo_datos.png"
    make_legacy_diagram(legacy_img)
    make_new_diagram(new_img)
    make_data_diagram(data_img)

    doc = Document()
    setup_styles(doc)
    add_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title.add_run("Cuadre automatico: flujo original legacy y flujo nuevo web")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(11, 37, 69)
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.add_paragraph()
    subtitle.add_run("Documento tecnico de referencia - generado el 04/08/2026").italic = True
    subtitle.paragraph_format.space_after = Pt(12)

    add_note(
        doc,
        "Objetivo",
        "Explicar como funciona el cuadre automatico original del sistema legacy y como esta quedando implementado en el nuevo front/back, incluyendo validaciones, tablas, SP y puntos de control.",
        LIGHT_BLUE,
    )

    doc.add_heading("1. Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "El cuadre automatico toma las rutas pendientes de cuadre para una fecha, calcula por producto el saldo inicial, lo usado en las OT del dia y los retiros, valida que las cantidades sigan coincidiendo y registra el cuadre si no hay bloqueos."
    )
    add_bullets(
        doc,
        [
            "El flujo original vive en una aplicacion WinForms legacy y ejecuta la logica ruta por ruta dentro del formulario.",
            "El flujo nuevo expone una pantalla para el usuario sistemas y delega la ejecucion al backend Spring mediante endpoints REST.",
            "La regla critica de materiales usados es contar solo filas activas: tbl_venta.e_eliminado = 0 y tbl_codigoventa.e_eliminado = 0.",
            "El nuevo flujo registra en transaccion: cabecera de cuadre, detalle por producto, retiros, actualizacion de saldo y snapshot de productos saldo.",
        ],
    )

    doc.add_heading("2. Alcance y fuentes del flujo", level=1)
    add_table(
        doc,
        ["Elemento", "Original legacy", "Nuevo sistema"],
        [
            ["Interfaz", "Formulario WinForms CuadresAutomatico.cs", "Pantalla SistemasCuadreAutomaticoPage.tsx"],
            ["Entrada principal", "Fecha y rutas no cuadradas", "Fecha, sucursal y token de usuario sistemas"],
            ["Ejecucion", "Proceso local del formulario", "POST /cuadre/sistemas/automatico/ejecutar"],
            ["Preview", "La vista prepara/valida dentro del flujo legacy", "GET /cuadre/sistemas/automatico/preview lista pendientes sin validar pesado"],
            ["Base de pruebas", "Base operativa configurada en legacy", "bdprueba para pruebas del nuevo flujo"],
        ],
        [1.5, 2.45, 2.45],
    )

    doc.add_heading("3. Flujo original legacy", level=1)
    doc.add_picture(str(legacy_img), width=Inches(6.35))
    cap = doc.add_paragraph("Figura 1. Secuencia general del cuadre automatico original.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)

    doc.add_heading("3.1 Pasos funcionales", level=2)
    add_numbered(
        doc,
        [
            "El operador ejecuta el formulario legacy de cuadre automatico y selecciona la fecha.",
            "El sistema obtiene las rutas pendientes con sp_TraerRutasNoCuadradas.",
            "Para cada ruta prepara datos: saldo actual de la ruta, catalogo de productos, productos retirados/no entregados y ventas del dia.",
            "Calcula detalle por producto: ItemsSobrantes, ItemsVendidos, ItemsRetirados, precio y total vendido.",
            "Ejecuta validaciones de cuadre existente, movimientos pendientes, consistencia de ventas y consistencia de saldo.",
            "Si la ruta pasa validacion, registra cabecera y detalle de cuadre; luego registra retiros y actualiza ruta/vendedor.",
            "Si falla, la ruta se marca como omitida o error y el proceso continua con la siguiente ruta.",
        ],
    )

    doc.add_heading("3.2 Preparacion de datos", level=2)
    add_table(
        doc,
        ["Dato", "Origen legacy", "Uso en el calculo"],
        [
            ["Saldo por producto", "TraerSaldoTarjetasRuta(idRuta)", "Base de ItemsSobrantes antes de descontar ventas"],
            ["Catalogo de productos", "TraerTodosLosProductos()", "Nombre y precio de producto"],
            ["Retiros/no entregados", "spx_ObtenerProductosNoEntregadosOT(idRuta)", "Carga ItemsRetirados y detalle tbl_SaldoRetiro"],
            ["Ventas del dia", "sp_TraerVentaDiaRuta(idRuta, fecha)", "Carga ItemsVendidos y descuenta del sobrante"],
            ["Cantidad de OT", "sp_TraerVentaDiaRuta_CantOt(idRuta, fecha)", "Control de que no cambie la cantidad de ordenes mientras se procesa"],
        ],
        [1.65, 2.35, 2.4],
    )

    doc.add_heading("3.3 Validaciones originales", level=2)
    add_bullets(
        doc,
        [
            "datos.Cargado debe ser verdadero; si no se pudo preparar la ruta, se omite.",
            "spx_ValidarCuadreRuta no debe encontrar un cuadre ya registrado para la ruta y fecha.",
            "ItemsVendidosVentas vuelve a leer ventas y cantidad de OT para detectar cambios durante la ejecucion.",
            "spx_ValidaMovimientos no debe reportar movimientos pendientes antes/despues del ultimo cierre.",
            "El detalle debe tener datos y la ruta debe ser valida.",
            "No se permite sobrante negativo.",
            "Para cada producto: ItemsSobrantes + ItemsVendidos debe coincidir con el saldo base.",
            "La cantidad vendida del detalle debe coincidir con sp_TraerVentaDiaRuta.",
        ],
    )

    doc.add_heading("4. Flujo nuevo web/back", level=1)
    doc.add_picture(str(new_img), width=Inches(6.35))
    cap = doc.add_paragraph("Figura 2. Secuencia general del flujo nuevo para usuario sistemas.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)

    doc.add_heading("4.1 Entrada por front", level=2)
    doc.add_paragraph(
        "El usuario sistemas entra a la pantalla de cuadre automatico, elige fecha y sucursal, revisa un preview de rutas pendientes y ejecuta el proceso. La pantalla no hace el calculo pesado; consume el resumen que devuelve el backend."
    )
    add_table(
        doc,
        ["Endpoint", "Funcion", "Observacion"],
        [
            ["GET /cuadre/sistemas/automatico/preview", "Lista rutas pendientes", "Usa sp_TraerRutasNoCuadradas; evita dejar la pantalla cargando por validaciones pesadas."],
            ["POST /cuadre/sistemas/automatico/ejecutar", "Ejecuta cuadre ruta por ruta", "Devuelve estado por ruta: Registrado, Omitido o Error."],
            ["/sistemas/cuadre-automatico", "Pantalla web", "Disponible solo para el usuario/rol sistemas."],
        ],
        [2.0, 2.0, 2.4],
    )

    doc.add_heading("4.2 Ejecucion backend", level=2)
    add_numbered(
        doc,
        [
            "requireSistemas valida que el token pertenezca al usuario/rol sistemas.",
            "resolverSucursalSistemas define la sucursal operativa donde se ejecuta el cuadre.",
            "resolverUsuarioRegistroAutomatico define el id_usuario que quedara registrado en tbl_Cuadre.",
            "obtenerRutasNoCuadradas trae las rutas pendientes para la fecha.",
            "crearResultadoRutaAutomatico procesa cada ruta con try/catch individual.",
            "validarRegistroPermitido bloquea cuadre duplicado, cierre de almacen, cierre PR/PD y movimientos pendientes.",
            "obtenerSaldoRutaConFallback calcula saldo/usado/sobrante; si falla, usa spb_SaldoRutasCantidad_X_Ruta.",
            "obtenerRetirosConFallback trae retiros con spx_ObtenerProductosNoEntregadosOT.",
            "normalizarDetalle arma el detalle final por producto.",
            "validarDetalleCuadre verifica consistencia antes de abrir la transaccion de registro.",
            "registrarCuadreTecnico ejecuta la transaccion SQL y confirma o revierte todo.",
        ],
    )

    doc.add_heading("4.3 Modelo de datos y formula", level=2)
    doc.add_picture(str(data_img), width=Inches(6.35))
    cap = doc.add_paragraph("Figura 3. Modelo de datos usado por el calculo nuevo.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)

    add_note(
        doc,
        "Regla de saldo usado hoy",
        "El usado del dia se calcula desde tbl_venta INNER JOIN tbl_codigoventa por fecha/ruta/producto, con cv.id_tipomaterial = 1, v.e_eliminado = 0 y cv.e_eliminado = 0. Si cualquiera de las dos tablas marca eliminado, esa cantidad no debe contarse.",
        RED,
    )

    doc.add_heading("5. Registro transaccional nuevo", level=1)
    add_table(
        doc,
        ["Paso", "SP / accion", "Tabla afectada"],
        [
            ["1", "spx_RegistrarCuadre", "tbl_Cuadre"],
            ["2", "spx_RegistrarCodigoCuadre por producto", "tbl_CodigoCuadre"],
            ["3", "spx_ActualizarSaldoTarjetaCuadre por producto", "tbl_saldotarjetas"],
            ["4", "spx_RegistrarSaldoRetiro por retiro", "tbl_SaldoRetiro"],
            ["5", "spx_RegistrarProductosSaldos 3, idRuta", "tbl_productosSaldos"],
            ["6", "sp_actualizarEquipoTecnico", "Equipo/ruta tecnica"],
        ],
        [0.65, 3.05, 2.7],
    )
    doc.add_paragraph(
        "La transaccion se confirma al final. Si cualquier SP o validacion SQL falla durante el registro, el backend ejecuta rollback y la ruta queda como Error sin dejar registros parciales de cuadre."
    )

    doc.add_heading("6. Diferencias principales", level=1)
    add_table(
        doc,
        ["Tema", "Original legacy", "Nuevo sistema"],
        [
            ["Control de acceso", "Depende de acceso al ejecutable/usuario legacy", "Solo usuario/rol sistemas por token"],
            ["Preview", "No separado claramente de ejecucion", "Preview rapido: lista rutas pendientes sin validar pesado"],
            ["Resultado por ruta", "Resumen textual del proceso", "Estado estructurado por ruta con contadores"],
            ["Validaciones", "Dentro del formulario WinForms", "Centralizadas en CuadreService"],
            ["Registro", "Componentes C# llaman capa CAD/CRN", "Repository Java llama SP dentro de una transaccion"],
            ["Base de pruebas", "Segun conexion legacy", "Debe apuntar a bdprueba durante pruebas"],
            ["Cierres automaticos", "El legacy tiene rutinas relacionadas a cierres", "El nuevo flujo por ahora valida cierres/movimientos; no genera cierres automaticos"],
        ],
        [1.45, 2.4, 2.55],
    )

    doc.add_heading("7. Casos de resultado", level=1)
    add_table(
        doc,
        ["Estado", "Cuando ocurre", "Accion esperada"],
        [
            ["Registrado", "La ruta paso validaciones y se confirmo la transaccion", "Mostrar idCuadre y sumar al contador Registradas"],
            ["Omitido", "Fallo una regla de negocio conocida: cierre, cuadre previo, saldo no coincide, sobrante negativo", "Mostrar mensaje y permitir correccion de datos"],
            ["Error", "Fallo tecnico: SP faltante, SQL exception, problema de conexion", "Revisar backend/base y volver a ejecutar"],
            ["Pendiente/Listo", "Preview o validacion sin registro", "No modifica tablas de cuadre"],
        ],
        [1.15, 3.1, 2.15],
    )

    doc.add_heading("8. Puntos criticos de operacion", level=1)
    add_bullets(
        doc,
        [
            "Durante pruebas se debe usar bdprueba; si el backend apunta a otra base, los saldos y eliminados no van a coincidir con lo que se revisa manualmente.",
            "Si una OT como 686 esta activa en tbl_venta pero eliminada en tbl_codigoventa, no debe sumar en usado hoy.",
            "El preview no garantiza que una ruta se registrara; la validacion real ocurre durante Ejecutar.",
            "Si se crean o cambian SP, deben existir en la misma base que usa la sucursal seleccionada.",
            "No conviene ejecutar cuadre mientras tecnicos siguen cargando OT/materiales, porque las validaciones comparan ventas contra detalle del momento.",
        ],
    )

    doc.add_heading("9. Checklist de prueba recomendado", level=1)
    add_numbered(
        doc,
        [
            "Confirmar que application.properties apunta a bdprueba o que la sucursal seleccionada resuelve bdprueba.",
            "Confirmar que existen los SP de registro: spx_RegistrarCuadre, spx_RegistrarCodigoCuadre, spx_ActualizarSaldoTarjetaCuadre, spx_RegistrarSaldoRetiro.",
            "Ejecutar preview y verificar cantidad de rutas pendientes.",
            "Ejecutar cuadre automatico en una fecha controlada.",
            "Revisar rutas omitidas por SALDO_NO_COINCIDE o VENTA_NO_COINCIDE antes de reintentar.",
            "Validar una ruta registrada contra tbl_Cuadre, tbl_CodigoCuadre, tbl_saldotarjetas y tbl_SaldoRetiro.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("10. Referencias de codigo revisadas", level=1)
    add_table(
        doc,
        ["Archivo", "Seccion relevante"],
        [
            ["CuadresAutomatico.cs", "PrepararDatosCuadre, ValidarCuadreAutomatico, RegistrarCuadreAutomatico"],
            ["CuadreService.java", "previewCuadreAutomaticoSistemas, ejecutarCuadreAutomaticoSistemas, validarRegistroPermitido, validarDetalleCuadre"],
            ["OtRepository.java", "obtenerSaldoRuta, obtenerRutasNoCuadradas, registrarCuadreTecnico y SP de registro"],
            ["SistemasCuadreAutomaticoPage.tsx", "Pantalla de parametros, preview, ejecucion y tabla de resultados"],
        ],
        [2.35, 4.05],
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    path = build_doc()
    print(path)
